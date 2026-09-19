"""Проверка содержимого сгенерированных PDF и XLSX (текст, кириллица, диакритика, суммы)."""
import glob
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402

OUT = Path("../docs/generated")
print("=" * 78)
print("PDF")
print("=" * 78)
for f in sorted(OUT.glob("*.pdf")):
    r = PdfReader(str(f))
    txt = " ".join((p.extract_text() or "") for p in r.pages)
    flat = " ".join(txt.split())
    has_cyr = any("а" <= ch.lower() <= "я" for ch in flat)
    has_dia = any(ch in flat for ch in "ăâîșțĂÂÎȘȚ")
    print(f"\n{f.name}  [{f.stat().st_size} Б, страниц {len(r.pages)}]")
    print(f"  кириллица: {'да' if has_cyr else 'НЕТ'} | диакритика: {'да' if has_dia else 'нет'}")
    print(f"  {flat[:330]}")

print("\n" + "=" * 78)
print("XLSX")
print("=" * 78)
for f in sorted(OUT.glob("*.xlsx")):
    wb = load_workbook(str(f))
    print(f"\n{f.name}  [{f.stat().st_size} Б]")
    for ws in wb.worksheets:
        head = [c.value for c in ws[1]][:6]
        print(f"  лист «{ws.title}»: строк {ws.max_row}, колонок {ws.max_column}")
        print(f"     шапка: {head}")
        if ws.max_row > 1:
            row2 = [c.value for c in ws[2]][:6]
            print(f"     пример: {row2}")

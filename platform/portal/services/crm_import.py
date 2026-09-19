"""Импорт данных из внешней CRM: CSV и Excel.

Модуль не привязан к конкретной системе: любая CRM умеет выгружать контакты, компании и сделки
в CSV/XLSX. Колонки распознаются по синонимам заголовков (RU/RO/EN), сопоставление можно
поправить вручную перед импортом. Дубликаты определяются по нормализованному телефону и e-mail.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import load_workbook

from ..db import SessionLocal as db
from ..models import Contact, Order, Partner, Tenant
from .orders import next_number, normalize_phone

# ---------------------------------------------------------------- распознавание колонок

FIELDS = {
    "name": ["имя", "фио", "контакт", "контактное лицо", "клиент", "name", "full name", "contact", "client",
             "nume", "persoana", "persoană de contact"],
    "company": ["компания", "организация", "фирма", "юрлицо", "company", "organization", "account",
                "companie", "firma", "firmă"],
    "phone": ["телефон", "тел", "мобильный", "phone", "mobile", "tel", "telefon", "număr"],
    "email": ["почта", "e-mail", "email", "мейл", "mail", "poșta"],
    "kind": ["тип", "тип клиента", "type", "segment", "сегмент", "tip"],
    "source": ["источник", "source", "канал", "channel", "sursa", "sursă"],
    "tags": ["теги", "метки", "tags", "labels", "etichete"],
    "notes": ["заметки", "комментарий", "описание", "notes", "comment", "description", "note", "observatii"],
    "idno": ["idno", "инн", "фискальный код", "код", "cod fiscal", "vat"],
    "city": ["город", "city", "oraș", "localitate"],
    "address": ["адрес", "address", "adresa", "adresă"],
    "amount": ["сумма", "оборот", "выручка", "amount", "revenue", "total", "suma", "valoare"],
    "orders_count": ["заказов", "сделок", "deals", "orders", "comenzi"],
    "created": ["создан", "дата", "created", "date", "data", "дата создания"],
    "manager": ["менеджер", "ответственный", "owner", "manager", "responsabil"],
    "external_id": ["id", "код записи", "external id", "crm id", "идентификатор"],
}
SYNONYM = {syn: field for field, syns in FIELDS.items() for syn in syns}


def norm_header(h: str) -> str:
    return re.sub(r"[^\w\s@-]", " ", str(h or "").strip().lower()).strip()


def detect_mapping(headers: list[str]) -> dict[int, str]:
    """Индекс колонки -> поле платформы."""
    mapping: dict[int, str] = {}
    used = set()
    for i, h in enumerate(headers):
        key = norm_header(h)
        field_name = SYNONYM.get(key)
        if not field_name:
            for syn, f in SYNONYM.items():
                if key and (key.startswith(syn) or syn in key):
                    field_name = f
                    break
        if field_name and field_name not in used:
            mapping[i] = field_name
            used.add(field_name)
    return mapping


# ---------------------------------------------------------------- чтение файлов

def read_table(data: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Возвращает заголовки и строки. Поддерживает CSV (utf-8/cp1251, ; , tab) и XLSX."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else (c.strftime("%d.%m.%Y") if hasattr(c, "strftime") else str(c)) for c in row])
        wb.close()
        if not rows:
            return [], []
        return rows[0], [r for r in rows[1:] if any(str(x).strip() for x in r)]

    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("не удалось определить кодировку файла")
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


# ---------------------------------------------------------------- импорт

@dataclass
class ImportReport:
    total: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    orders: int = 0
    errors: list[str] = field(default_factory=list)
    samples: list[dict] = field(default_factory=list)
    dry_run: bool = True

    @property
    def ok(self) -> bool:
        return not self.errors

    def line(self) -> str:
        mode = "пробный прогон" if self.dry_run else "импорт"
        return (f"{mode}: строк {self.total}, создано {self.created}, обновлено {self.updated}, "
                f"пропущено {self.skipped}" + (f", заявок {self.orders}" if self.orders else ""))


_IDNO_RE = re.compile(r"IDNO:\s*([\d]{6,20})")
_EXT_RE = re.compile(r"ID в CRM:\s*(\S+)")
_COMPANY_NOISE = re.compile(r"(?i)\b(srl|sa|s\.?r\.?l\.?|s\.?a\.?|ооо|оао|зао|ип|î\.?i\.?|gmbh|ltd|llc)\b|[^\w\s]")


def _company_key(name: str) -> str:
    """Название компании без организационной формы и знаков — для сопоставления выгрузок."""
    s = _COMPANY_NOISE.sub(" ", (name or "").lower())
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) >= 3 else ""


def _idno_of(contact) -> str:
    m = _IDNO_RE.search(contact.notes or "")
    return m.group(1) if m else ""


def _external_of(contact) -> str:
    m = _EXT_RE.search(contact.notes or "")
    return m.group(1) if m else ""


def parse_amount(v) -> float:
    s = re.sub(r"[^\d,.\-]", "", str(v or "")).replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def import_contacts(tenant: Tenant, headers: list[str], rows: list[list[str]], mapping: dict[int, str],
                    *, dry_run: bool = True, source_label: str = "crm-import",
                    create_orders: bool = False, partner_id: int | None = None) -> ImportReport:
    rep = ImportReport(dry_run=dry_run)
    existing = db.query(Contact).filter_by(tenant_id=tenant.id).all()
    by_phone = {c.phone: c for c in existing if c.phone}
    by_email = {c.email.lower(): c for c in existing if c.email}
    # выгрузки ERP часто идут без телефона и почты: там опознаём по фискальному коду,
    # внешнему идентификатору и нормализованному названию компании
    by_idno = {k: c for c in existing for k in [_idno_of(c)] if k}
    by_external = {k: c for c in existing for k in [_external_of(c)] if k}
    by_company = {k: c for c in existing for k in [_company_key(c.company)] if k}

    for n, row in enumerate(rows, start=2):
        try:
            rec = {}
            for idx, fld in mapping.items():
                if idx < len(row):
                    rec[fld] = str(row[idx]).strip()
            name = rec.get("name") or rec.get("company")
            phone = normalize_phone(rec.get("phone", ""))
            email = (rec.get("email") or "").strip().lower()
            if not (name or phone or email):
                rep.skipped += 1
                continue
            rep.total += 1

            idno = re.sub(r"\D", "", rec.get("idno", ""))
            external = (rec.get("external_id") or "").strip()
            company_key = _company_key(rec.get("company", ""))
            found = ((by_phone.get(phone) if phone else None)
                     or (by_email.get(email) if email else None)
                     or (by_idno.get(idno) if idno else None)
                     or (by_external.get(external) if external else None)
                     or (by_company.get(company_key) if company_key and not (phone or email) else None))
            payload = {
                "name": name or phone or email,
                "company": rec.get("company", ""),
                "phone": phone,
                "email": email,
                "kind": "b2b" if (rec.get("company") or rec.get("idno") or partner_id) else "b2c",
                "source": source_label,
                "tags": rec.get("tags", ""),
                "notes": "\n".join(x for x in [rec.get("notes", ""),
                                               f"IDNO: {rec['idno']}" if rec.get("idno") else "",
                                               f"Город: {rec['city']}" if rec.get("city") else "",
                                               f"Адрес: {rec['address']}" if rec.get("address") else "",
                                               f"Менеджер в CRM: {rec['manager']}" if rec.get("manager") else "",
                                               f"ID в CRM: {rec['external_id']}" if rec.get("external_id") else ""] if x),
                "revenue": parse_amount(rec.get("amount")),
                "orders_count": int(parse_amount(rec.get("orders_count")) or 0),
            }
            if len(rep.samples) < 8:
                rep.samples.append({k: v for k, v in payload.items() if v})

            if dry_run:
                rep.updated += 1 if found else 0
                rep.created += 0 if found else 1
                continue

            if found:
                for k, v in payload.items():
                    if v and not getattr(found, k, None):
                        setattr(found, k, v)
                contact = found
                rep.updated += 1
            else:
                contact = Contact(tenant_id=tenant.id, partner_id=partner_id, **payload)
                db.add(contact)
                db.flush()
                if phone:
                    by_phone[phone] = contact
                if email:
                    by_email[email] = contact
                rep.created += 1

            if create_orders and payload["revenue"]:
                db.add(Order(tenant_id=tenant.id, contact_id=contact.id, number=next_number(tenant),
                             kind=payload["kind"], source=source_label, stage="won", status="completed",
                             partner_id=partner_id, task_type="imported", cargo="перенос истории из CRM",
                             price_final=payload["revenue"], starts_at=datetime.utcnow(),
                             comment=f"Импортировано из внешней CRM, строка {n}"))
                rep.orders += 1
        except Exception as exc:  # noqa: BLE001
            rep.errors.append(f"строка {n}: {exc}")
            rep.skipped += 1
    if not dry_run:
        db.commit()
    return rep


def preview(headers: list[str], rows: list[list[str]], mapping: dict[int, str], limit: int = 8) -> list[dict]:
    out = []
    for row in rows[:limit]:
        rec = {}
        for idx, fld in mapping.items():
            if idx < len(row) and str(row[idx]).strip():
                rec[fld] = str(row[idx]).strip()
        if rec:
            out.append(rec)
    return out


FIELD_TITLES = {
    "name": "Имя", "company": "Компания", "phone": "Телефон", "email": "E-mail", "kind": "Тип",
    "source": "Источник", "tags": "Теги", "notes": "Заметки", "idno": "IDNO / фискальный код",
    "city": "Город", "address": "Адрес", "amount": "Оборот", "orders_count": "Число заказов",
    "created": "Дата создания", "manager": "Менеджер", "external_id": "ID в CRM",
}

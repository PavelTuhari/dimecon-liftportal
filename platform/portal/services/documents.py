"""Генерация документов: PDF (коммерческое предложение, подтверждение, акт, счёт, рапорт)
и выгрузки в Excel (заявки, смены, счета, парк, прайс, контакты).

PDF — fpdf2 со шрифтом DejaVu (кириллица + румынская диакритика ș/ț).
XLSX — openpyxl, без внешних зависимостей.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_DIR = Path(__file__).resolve().parent.parent.parent / "assets" / "fonts"
ACCENT = (245, 166, 35)
NAVY = (31, 58, 82)
GREY = (90, 100, 107)
LINE = (221, 225, 228)

DOC_TITLES = {
    "quote": "Коммерческое предложение",
    "confirmation": "Подтверждение заказа",
    "act": "Акт выполненных работ",
    "invoice": "Счёт на оплату",
    "shift_report": "Сменный рапорт",
}

_ONES = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять", "десять",
         "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать", "шестнадцать",
         "семнадцать", "восемнадцать", "девятнадцать"]
_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
_HUNDREDS = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]


def _triple(n: int, female: bool) -> list[str]:
    out = []
    if n >= 100:
        out.append(_HUNDREDS[n // 100])
        n %= 100
    if n >= 20:
        out.append(_TENS[n // 10])
        n %= 10
    if n:
        word = _ONES[n]
        if female and n == 1:
            word = "одна"
        elif female and n == 2:
            word = "две"
        out.append(word)
    return out


# склонение названий валют: (1, 2-4, 5+) для целой и дробной части
CURRENCY_WORDS = {
    "MDL": (("лей", "лея", "леев"), ("бан", "бана", "бань")),
    "RON": (("лей", "лея", "леев"), ("бан", "бана", "бань")),
    "EUR": (("евро", "евро", "евро"), ("цент", "цента", "центов")),
    "USD": (("доллар", "доллара", "долларов"), ("цент", "цента", "центов")),
    "RUB": (("рубль", "рубля", "рублей"), ("копейка", "копейки", "копеек")),
}


def plural(n: int, forms: tuple[str, str, str]) -> str:
    last, tail = n % 10, n % 100
    if 11 <= tail <= 14 or last == 0 or last >= 5:
        return forms[2]
    return forms[0] if last == 1 else forms[1]


def amount_in_words(value: float, currency: str = "MDL") -> str:
    """Сумма прописью для счёта и акта: «Шестьдесят две тысячи четыреста восемьдесят леев 00 бань»."""
    total = int(round(value * 100))
    whole, cents = divmod(total, 100)
    if whole == 0:
        words = ["ноль"]
    else:
        words, groups = [], [(whole // 1_000_000, False, ("миллион", "миллиона", "миллионов")),
                             (whole // 1000 % 1000, True, ("тысяча", "тысячи", "тысяч")),
                             (whole % 1000, False, None)]
        for num, female, forms in groups:
            if not num:
                continue
            words += _triple(num, female)
            if forms:
                last, tail = num % 10, num % 100
                if 11 <= tail <= 14 or last == 0 or last >= 5:
                    words.append(forms[2])
                elif last == 1:
                    words.append(forms[0])
                else:
                    words.append(forms[1])
    text = " ".join(w for w in words if w)
    major, minor = CURRENCY_WORDS.get((currency or "MDL").upper(),
                                      ((currency, currency, currency), ("", "", "")))
    tail = f" {cents:02d} {plural(cents, minor)}".rstrip() if minor[0] else ""
    return f"{text[:1].upper()}{text[1:]} {plural(whole, major)}{tail}"


def money(v, cur="") -> str:
    try:
        s = f"{float(v):,.2f}".replace(",", " ").replace(".", ",")
    except (TypeError, ValueError):
        return str(v)
    return f"{s} {cur}".strip()


# --------------------------------------------------------------------------- PDF

class DocPDF(FPDF):
    def __init__(self, tenant, doc_title: str, doc_number: str = ""):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.tenant = tenant
        self.doc_title = doc_title
        self.doc_number = doc_number
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(15, 15, 15)
        self.add_font("dj", "", str(FONT_DIR / "DejaVuSans.ttf"))
        self.add_font("dj", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
        self.set_font("dj", "", 10)
        theme = tenant.theme or {}
        self.accent = _hex_rgb(theme.get("accent"), ACCENT)
        self.navy = _hex_rgb(theme.get("navy"), NAVY)
        self.add_page()

    def header(self):
        t = self.tenant
        self.set_fill_color(*self.navy)
        self.rect(0, 0, 210, 26, "F")
        self.set_fill_color(*self.accent)
        self.rect(0, 26, 210, 1.6, "F")
        self.set_xy(15, 7)
        self.set_text_color(255, 255, 255)
        self.set_font("dj", "B", 14)
        self.cell(110, 6, t.name, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_x(15)
        self.set_font("dj", "", 8)
        req = " · ".join(x for x in [t.legal_name, f"IDNO {t.idno}" if t.idno else "", t.address] if x)
        self.cell(110, 4.5, req[:95])
        self.set_xy(130, 7)
        self.set_font("dj", "", 8)
        for line in [t.phone, t.phone2, t.email]:
            if line:
                self.set_x(130)
                self.cell(65, 4.2, line, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_text_color(20, 24, 27)
        self.set_y(36)

    def footer(self):
        self.set_y(-15)
        self.set_draw_color(*LINE)
        self.line(15, self.get_y() - 2, 195, self.get_y() - 2)
        self.set_font("dj", "", 7.5)
        self.set_text_color(*GREY)
        self.cell(120, 5, f"{self.doc_title} {self.doc_number} · сформирован {datetime.now():%d.%m.%Y %H:%M}")
        self.cell(60, 5, f"стр. {self.page_no()} / {{nb}}", align="R")
        self.set_text_color(20, 24, 27)

    # --- строительные блоки ------------------------------------------------
    def title_block(self, subtitle: str = ""):
        self.set_font("dj", "B", 17)
        self.cell(0, 9, f"{self.doc_title} {self.doc_number}".strip(), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if subtitle:
            self.set_font("dj", "", 9.5)
            self.set_text_color(*GREY)
            self.multi_cell(0, 5, subtitle)
            self.set_text_color(20, 24, 27)
        self.ln(3)

    def kv_table(self, rows: list[tuple[str, str]], label_w=58):
        self.set_font("dj", "", 9.5)
        for k, v in rows:
            if v in (None, ""):
                continue
            y0 = self.get_y()
            self.set_text_color(*GREY)
            self.multi_cell(label_w, 5.6, str(k), align="L", new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.set_text_color(20, 24, 27)
            self.set_xy(15 + label_w, y0)
            self.multi_cell(180 - label_w, 5.6, str(v))
            self.set_draw_color(*LINE)
            self.line(15, self.get_y(), 195, self.get_y())
            self.ln(0.8)
        self.ln(2)

    def money_table(self, lines: list[dict], total: float, currency: str, total_label="ИТОГО к оплате"):
        self.set_font("dj", "B", 8.5)
        self.set_fill_color(240, 242, 243)
        self.cell(12, 7, "№", border=0, fill=True, align="C")
        self.cell(123, 7, "Наименование", border=0, fill=True)
        self.cell(45, 7, "Сумма", border=0, fill=True, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_font("dj", "", 9)
        for i, ln in enumerate(lines, 1):
            title = str(ln.get("title", ""))
            amount = ln.get("amount", 0)
            h = 6 if len(title) < 72 else 11
            y0 = self.get_y()
            self.cell(12, h, str(i), align="C")
            self.multi_cell(123, 5.5, title, new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.set_xy(150, y0)
            self.cell(45, h, money(amount), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_y(max(self.get_y(), y0 + h))
            self.set_draw_color(*LINE)
            self.line(15, self.get_y(), 195, self.get_y())
        self.ln(1)
        self.set_font("dj", "B", 11)
        self.set_fill_color(*self.accent)
        self.cell(135, 9, f"  {total_label}", fill=True)
        self.cell(45, 9, money(total, currency), align="R", fill=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(2)
        self.set_font("dj", "", 8.5)
        self.set_text_color(*GREY)
        self.multi_cell(0, 4.6, f"Сумма прописью: {amount_in_words(total, currency)}")
        self.set_text_color(20, 24, 27)
        self.ln(2)

    def note(self, text: str):
        self.set_font("dj", "", 8.5)
        self.set_fill_color(252, 239, 217)
        self.multi_cell(0, 4.8, text, fill=True)
        self.ln(2)

    def signatures(self, left="Исполнитель", right="Заказчик"):
        self.ln(6)
        y = self.get_y()
        self.set_font("dj", "", 9)
        for x, label in ((15, left), (110, right)):
            self.set_xy(x, y)
            self.cell(85, 5, label, new_x=XPos.LEFT, new_y=YPos.NEXT)
            self.set_x(x)
            self.set_draw_color(*GREY)
            self.line(x, self.get_y() + 6, x + 75, self.get_y() + 6)
            self.set_xy(x, self.get_y() + 7)
            self.set_font("dj", "", 7.5)
            self.set_text_color(*GREY)
            self.cell(85, 4, "подпись, Ф.И.О.")
            self.set_text_color(20, 24, 27)
            self.set_font("dj", "", 9)

    def out(self) -> bytes:
        return bytes(self.output())


def _hex_rgb(value, default):
    try:
        v = (value or "").lstrip("#")
        return int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16)
    except (ValueError, IndexError, AttributeError):
        return default


def _order_rows(order, tenant) -> list[tuple[str, str]]:
    eq = order.equipment
    return [
        ("Заявка", order.number),
        ("Дата подачи", f"{order.starts_at:%d.%m.%Y %H:%M}" if order.starts_at else "по согласованию"),
        ("Продолжительность", f"{order.hours:g} ч" if order.hours else ""),
        ("Объект", order.address or ""),
        ("Техника", f"{eq.brand} {eq.model}" if eq else "подбирается инженером"),
        ("Груз", f"{order.cargo} · {order.weight_t:g} т" if order.cargo else ""),
        ("Высота / вылет", f"{order.height_m:g} м / {order.radius_m:g} м" if order.height_m or order.radius_m else ""),
        ("Условия", ", ".join(order.conditions or []) or "обычные"),
    ]


def _client_rows(order) -> list[tuple[str, str]]:
    c = order.contact
    if not c:
        return []
    return [("Заказчик", c.company or c.name), ("Контактное лицо", c.name if c.company else ""),
            ("Телефон", c.phone), ("E-mail", c.email)]


def order_confirmation_pdf(tenant, order) -> bytes:
    pdf = DocPDF(tenant, DOC_TITLES["confirmation"], order.number)
    pdf.title_block("Заявка принята в работу. Ниже — согласованные параметры выезда техники.")
    pdf.kv_table(_client_rows(order) + _order_rows(order, tenant))
    if order.breakdown:
        total = order.price_final or order.price_max or sum(x.get("amount", 0) for x in order.breakdown)
        pdf.money_table(order.breakdown, total, tenant.currency, "Ориентировочно")
        pdf.note("Стоимость предварительная. Итоговая сумма определяется по сменному рапорту: "
                 "фактические часы, простои по вине заказчика и дополнительные услуги.")
    pdf.kv_table([("Диспетчер", tenant.phone), ("Электронная почта", tenant.email)])
    return pdf.out()


def quote_pdf(tenant, order, doc) -> bytes:
    pdf = DocPDF(tenant, DOC_TITLES["quote"], doc.number)
    valid = f"Предложение действительно до {doc.due_at:%d.%m.%Y}." if doc.due_at else ""
    pdf.title_block(f"К заявке {order.number}. {valid}")
    pdf.kv_table(_client_rows(order) + _order_rows(order, tenant))
    lines = order.breakdown or [{"title": "Услуги спецтехники по заявке " + order.number, "amount": doc.amount}]
    pdf.money_table(lines, doc.amount, tenant.currency)
    pdf.note("В стоимость включены: оператор, топливо, подача и возврат техники, страхование гражданской "
             "ответственности. Оплачивается отдельно: простой по вине заказчика, часы сверх заказанных, "
             "оформление разрешений.")
    pdf.signatures("Исполнитель — " + tenant.name, "Заказчик")
    return pdf.out()


def sales_pdf(tenant, so) -> bytes:
    """Коммерческое предложение из модуля продаж: строки сметы, скидка, НДС, допработы."""
    from ..db import SessionLocal as db
    from ..models import Contact, Partner

    pdf = DocPDF(tenant, "Коммерческое предложение", so.number)
    valid = f" Действительно до {so.valid_until:%d.%m.%Y}." if so.valid_until else ""
    pdf.title_block(f"от {so.created_at:%d.%m.%Y}.{valid}")

    client = ""
    if so.contact_id:
        c = db.get(Contact, so.contact_id)
        client = (c.company or c.name) if c else ""
    elif so.partner_id:
        p = db.get(Partner, so.partner_id)
        client = p.name if p else ""
    rows = [("Заказчик", client or "—"), ("Исполнитель", tenant.legal_name or tenant.name)]
    if tenant.idno:
        rows.append(("IDNO", tenant.idno))
    if so.discount_pct:
        rows.append(("Скидка", f"{so.discount_pct:g} %"))
    pdf.kv_table(rows)

    lines = [{"title": f"{l.name} ({l.qty:g} {l.unit} × {money(l.price)})", "amount": l.amount}
             for l in so.lines if not l.is_optional]
    if so.discount_pct:
        lines.append({"title": f"Скидка {so.discount_pct:g} %",
                      "amount": -round(sum(l["amount"] for l in lines) * so.discount_pct / 100, 2)})
    lines.append({"title": f"НДС {so.vat_pct:g} %", "amount": so.amount_vat})
    pdf.money_table(lines, so.amount_total, so.currency)

    optional = [l for l in so.lines if l.is_optional]
    if optional:
        pdf.note("Дополнительно, по желанию заказчика: "
                 + "; ".join(f"{l.name} — {money(l.amount, so.currency)}" for l in optional))
    if so.terms:
        pdf.note(so.terms)
    pdf.note(f"Сумма прописью: {amount_in_words(so.amount_total, so.currency)}.")
    pdf.signatures("Исполнитель — " + tenant.name, "Заказчик")
    return pdf.out()


def invoice_pdf(tenant, order, doc) -> bytes:
    pdf = DocPDF(tenant, DOC_TITLES["invoice"], doc.number)
    pdf.title_block(f"от {doc.issued_at:%d.%m.%Y} по заявке {order.number}"
                    + (f". Срок оплаты до {doc.due_at:%d.%m.%Y}." if doc.due_at else ""))
    bank = tenant.settings.get("bank") if tenant.settings else None
    pdf.kv_table(_client_rows(order) + [("Поставщик", tenant.legal_name or tenant.name),
                                        ("IDNO", tenant.idno), ("Банковские реквизиты", bank or "по договору")])
    lines = order.breakdown or [{"title": f"Услуги спецтехники по заявке {order.number}", "amount": doc.amount}]
    pdf.money_table(lines, doc.amount, tenant.currency)
    pdf.note("Оплатой счёта заказчик подтверждает согласие с условиями аренды, включая порядок отмены: "
             "более чем за 24 часа — бесплатно, за 6–24 часа — 30 % минимальной смены, менее 6 часов — "
             "полная минимальная смена и подача.")
    pdf.signatures("Руководитель / гл. бухгалтер", "Принял")
    return pdf.out()


def act_pdf(tenant, order, doc, shifts=()) -> bytes:
    pdf = DocPDF(tenant, DOC_TITLES["act"], doc.number)
    pdf.title_block(f"от {doc.issued_at:%d.%m.%Y} по заявке {order.number}")
    pdf.kv_table(_client_rows(order) + _order_rows(order, tenant))
    if shifts:
        pdf.set_font("dj", "B", 10)
        pdf.cell(0, 7, "Выполненные смены", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("dj", "B", 8.5)
        pdf.set_fill_color(240, 242, 243)
        for w, h in ((24, "Дата"), (46, "Оператор"), (24, "Часы"), (26, "Простой"), (24, "Подъёмов"), (36, "Статус")):
            pdf.cell(w, 7, h, fill=True)
        pdf.ln()
        pdf.set_font("dj", "", 9)
        for s in shifts:
            pdf.cell(24, 6, f"{s.work_date:%d.%m.%Y}")
            pdf.cell(46, 6, (s.operator or "—")[:24])
            pdf.cell(24, 6, f"{s.hours_worked:g}")
            pdf.cell(26, 6, f"{s.hours_idle:g}" + (f" ({s.idle_fault})" if s.idle_fault else ""))
            pdf.cell(24, 6, str(s.lifts or 0))
            pdf.cell(36, 6, s.status)
            pdf.ln()
            pdf.set_draw_color(*LINE)
            pdf.line(15, pdf.get_y(), 195, pdf.get_y())
        pdf.ln(3)
    lines = order.breakdown or [{"title": f"Услуги спецтехники по заявке {order.number}", "amount": doc.amount}]
    pdf.money_table(lines, doc.amount, tenant.currency, "ИТОГО выполнено")
    pdf.note("Работы выполнены в полном объёме. Претензий по объёму, качеству и срокам стороны не имеют. "
             "Акт составлен на основании подтверждённых сменных рапортов.")
    pdf.signatures("Исполнитель — " + tenant.name, "Заказчик")
    return pdf.out()


def shift_report_pdf(tenant, shift, order=None) -> bytes:
    pdf = DocPDF(tenant, DOC_TITLES["shift_report"], f"№ {shift.id} от {shift.work_date:%d.%m.%Y}")
    pdf.title_block("Первичный документ учёта работы техники на объекте.")
    # у модели смены нет связи с техникой — берём её из заявки
    eq = getattr(order, "equipment", None) if order else None
    rows = [("Заявка", order.number if order else str(shift.order_id)),
            ("Объект", order.address if order else ""),
            ("Техника", f"{eq.brand} {eq.model}".strip() if eq else "не назначена"),
            ("Оператор", shift.operator or ""),
            ("Часы работы", f"{shift.hours_worked:g} ч"),
            ("Простой", f"{shift.hours_idle:g} ч" + (f" (вина: {shift.idle_fault})" if shift.idle_fault else "")),
            ("Количество подъёмов", str(shift.lifts or 0)),
            ("Статус рапорта", shift.status),
            ("Примечания", shift.notes or "—")]
    if shift.dispute_reason:
        rows.append(("Основание спора", shift.dispute_reason))
    pdf.kv_table(rows)
    pdf.note("Подтверждённый рапорт является основанием для акта выполненных работ и счёта. "
             "Возражения принимаются в течение 72 часов с момента передачи.")
    pdf.signatures("Оператор", "Представитель заказчика")
    return pdf.out()


def document_pdf(tenant, order, doc, shifts=()) -> bytes:
    if doc.type == "quote":
        return quote_pdf(tenant, order, doc)
    if doc.type == "invoice":
        return invoice_pdf(tenant, order, doc)
    if doc.type == "act":
        return act_pdf(tenant, order, doc, shifts)
    return order_confirmation_pdf(tenant, order)


# --------------------------------------------------------------------------- XLSX

HEAD_FILL = PatternFill("solid", fgColor="1F3A52")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="DDE1E4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb, title: str, headers: list[str], widths: list[int]):
    ws = wb.active if wb.active.max_row == 1 and wb.active.max_column == 1 and wb.active.title == "Sheet" else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    return ws


def _finish(wb) -> bytes:
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(str(cell.value)) > 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def orders_xlsx(tenant, orders, title="Заявки") -> bytes:
    wb = Workbook()
    ws = _sheet(wb, title, ["№ заявки", "Создана", "Тип", "Источник", "Клиент", "Телефон", "Компания",
                            "Задача", "Груз", "Масса, т", "Высота, м", "Вылет, м", "Техника", "Подача",
                            "Часы", "Адрес", "Зона", "Этап", "Статус", f"Оценка, {tenant.currency}",
                            f"Итог, {tenant.currency}", "Эскалация"],
                [16, 16, 7, 12, 22, 16, 22, 14, 24, 10, 11, 10, 22, 16, 8, 30, 10, 12, 14, 16, 14, 30])
    for o in orders:
        c = o.contact
        ws.append([o.number, o.created_at, o.kind.upper(), o.source, c.name if c else "", c.phone if c else "",
                   c.company if c else "", o.task_type, o.cargo, o.weight_t, o.height_m, o.radius_m,
                   f"{o.equipment.brand} {o.equipment.model}" if o.equipment else "",
                   o.starts_at, o.hours, o.address, o.zone, o.stage, o.status,
                   o.price_max or o.price_min or 0, o.price_final or 0,
                   "; ".join(o.escalation_reasons or [])])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = row[13].number_format = "DD.MM.YYYY HH:MM"
        for i in (19, 20):
            row[i].number_format = "# ##0.00"
    ws.append([])
    ws.append(["ИТОГО", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
               f"=SUM(T2:T{ws.max_row - 2})", f"=SUM(U2:U{ws.max_row - 2})", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    return _finish(wb)


def shifts_xlsx(tenant, shifts, orders_by_id=None) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Сменные рапорты", ["Дата", "Заявка", "Объект", "Техника", "Оператор", "Часы работы",
                                        "Простой, ч", "Вина простоя", "Подъёмов", "Статус", "Примечания"],
                [12, 16, 30, 22, 22, 12, 12, 14, 11, 14, 36])
    orders_by_id = orders_by_id or {}
    for s in shifts:
        o = orders_by_id.get(s.order_id)
        ws.append([s.work_date, o.number if o else s.order_id, o.address if o else "",
                   f"{o.equipment.brand} {o.equipment.model}" if o and o.equipment else "",
                   s.operator, s.hours_worked, s.hours_idle, s.idle_fault, s.lifts, s.status, s.notes])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
    n = ws.max_row
    ws.append(["ИТОГО", "", "", "", "", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", "", f"=SUM(I2:I{n})", "", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    return _finish(wb)


def invoices_xlsx(tenant, documents, orders_by_id=None) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Документы", ["Тип", "Номер", "Дата", "Срок оплаты", "Заявка", f"Сумма, {tenant.currency}",
                                  f"Оплачено, {tenant.currency}", "Остаток", "Статус"],
                [14, 18, 12, 14, 16, 16, 16, 14, 14])
    orders_by_id = orders_by_id or {}
    for d in documents:
        o = orders_by_id.get(d.order_id)
        # оплата фиксируется статусом документа: отдельного поля «оплачено» в модели нет
        paid = (d.amount or 0) if d.status == "paid" else 0
        ws.append([DOC_TITLES.get(d.type, d.type), d.number, d.issued_at, d.due_at,
                   o.number if o else "", d.amount, paid, (d.amount or 0) - paid, d.status])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = row[3].number_format = "DD.MM.YYYY"
        for i in (5, 6, 7):
            row[i].number_format = "# ##0.00"
    n = ws.max_row
    ws.append(["ИТОГО", "", "", "", "", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    return _finish(wb)


def fleet_xlsx(tenant, equipment) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Парк техники", ["Инв. №", "Категория", "Бренд", "Модель", "Год", "Г/П, т", "Вылет, м",
                                     "Высота, м", "Грузоподъёмность, кг", f"Ставка/ч, {tenant.currency}",
                                     "Мин. смена, ч", f"Подача, {tenant.currency}", "Строк грузовой таблицы",
                                     "Статус", "Опубликовано"],
                [10, 22, 14, 22, 7, 10, 10, 10, 18, 16, 13, 14, 20, 12, 13])
    for e in equipment:
        ws.append([e.inventory_no, (e.category.name or {}).get("ru", "") if e.category else "", e.brand, e.model,
                   e.year, e.capacity_t or "", e.radius_m or "", e.height_m or "", e.payload_kg or "",
                   e.hourly_rate, e.min_hours, e.mobilization_fee, len(e.load_charts), e.status,
                   "да" if e.is_published else "нет"])
    ws2 = _sheet(wb, "Грузовые таблицы", ["Техника", "Конфигурация", "Опоры", "Вылет, м", "Г/П, т", "Высота, м"],
                 [28, 18, 12, 12, 12, 12])
    for e in equipment:
        for c in e.load_charts:
            ws2.append([f"{e.brand} {e.model}", c.configuration, c.outriggers, c.radius_m, c.capacity_t, c.height_m])
    return _finish(wb)


def pricelist_xlsx(tenant, equipment, services, products, pricing: dict) -> bytes:
    wb = Workbook()
    cur = tenant.currency
    ws = _sheet(wb, "Техника", ["Категория", "Техника", "Г/П, т", f"Ставка/ч, {cur}", "Мин. смена, ч",
                                f"Подача, {cur}", f"За км, {cur}", f"Минимальный заказ, {cur}"],
                [22, 26, 10, 16, 13, 14, 13, 20])
    for e in equipment:
        if not e.is_published:
            continue
        ws.append([(e.category.name or {}).get("ru", "") if e.category else "", f"{e.brand} {e.model}",
                   e.capacity_t or "", e.hourly_rate, e.min_hours, e.mobilization_fee, e.per_km_rate,
                   (e.hourly_rate or 0) * (e.min_hours or 0) + (e.mobilization_fee or 0)])
    ws1 = _sheet(wb, "Услуги", ["Услуга", "Модель цены", f"Цена от, {cur}"], [60, 16, 16])
    for s in services:
        if s.is_published:
            ws1.append([(s.title or {}).get("ru", ""), s.pricing_model, s.price_from])
    ws2 = _sheet(wb, "Товары", ["Артикул", "Категория", "Наименование", "Ед.", f"Цена, {cur}", "Остаток"],
                 [16, 24, 46, 8, 14, 10])
    for p in products:
        if p.is_published:
            ws2.append([p.sku, p.category, (p.title or {}).get("ru", ""), p.unit, p.price, p.stock])
    ws3 = _sheet(wb, "Коэффициенты", ["Группа", "Параметр", "Значение"], [22, 40, 18])
    for group in ("time", "season", "conditions", "addons", "discounts"):
        for k, v in (pricing.get(group) or {}).items():
            ws3.append([group, k, v])
    for z in pricing.get("zones", []):
        ws3.append(["zone", f"{z['name']} (до {z['max_km']} км)",
                    f"подача {z['fee']}, бесплатно {z['free_km']} км, далее {z['per_km']}/км"])
    return _finish(wb)


def contacts_xlsx(tenant, contacts) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Контакты", ["Имя", "Компания", "Телефон", "E-mail", "Тип", "Источник", "Заказов",
                                 f"Выручка, {tenant.currency}", "Теги", "Создан"],
                [26, 26, 18, 26, 8, 14, 10, 16, 20, 16])
    for c in contacts:
        ws.append([c.name, c.company, c.phone, c.email, c.kind.upper(), c.source, c.orders_count,
                   c.revenue, c.tags, c.created_at])
    for row in ws.iter_rows(min_row=2):
        row[9].number_format = "DD.MM.YYYY"
    return _finish(wb)
